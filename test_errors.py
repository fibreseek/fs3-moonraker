
import pexpect
import sys
import time
import json
import re
import shlex
import os
from typing import List, Tuple

class MoonrakerTester:
    def __init__(self, host, user, password):
        self.host = host
        self.user = user
        self.password = password
        self.config_path = "~/printer_data/config/moonraker.conf"
        self.log_path = "~/printer_data/logs/moonraker.log"
        self.xlsx_path = os.path.join(os.path.dirname(__file__), "error_info.xlsx")

    def log(self, msg):
        print(f"[TESTER] {msg}")

    def run_cmd(self, cmd, timeout=30, check_exit=True):
        # Use shlex.quote to avoid broken nested quoting when cmd contains quotes/braces
        ssh_cmd = (
            "ssh "
            "-o StrictHostKeyChecking=no "
            "-o UserKnownHostsFile=/dev/null "
            "-o LogLevel=ERROR "
            f"{self.user}@{self.host} "
            f"{shlex.quote(cmd)}"
        )
        child = pexpect.spawn(ssh_cmd)
        try:
            # Handle both 'password:' and 'Password:' prompts
            index = child.expect([r'[Pp]assword:', pexpect.EOF], timeout=timeout)
            if index == 0:
                child.sendline(self.password)
                child.expect(pexpect.EOF, timeout=timeout)
            
            output = child.before.decode() if child.before else ""
            child.close()
            
            if check_exit and child.exitstatus != 0:
                return False, output
            return True, output
        except Exception as e:
            return False, str(e)

    def fix_config_if_needed(self):
        self.log("Checking for duplicate config...")
        success, output = self.run_cmd(f"cat {self.config_path}")
        if not success:
            self.log(f"Failed to read config via SSH. Output:\n{output}")
            return False
        
        if output.count("[error_indexer]") > 1:
            self.log("⚠️ Duplicate config found! Fixing...")
            # Remove all [error_indexer] sections and re-add one
            # We use python to do this cleanly
            py_script = (
                "import re; "
                f"path = '{self.config_path}'.replace('~', '/home/{self.user}'); "
                "with open(path, 'r') as f: content = f.read(); "
                "content = re.sub(r'\\n\\[error_indexer\\][\\s\\S]*?(?=\\n\\[|$)', '', content); "
                "content += '\\n[error_indexer]\\nfile_path: /home/anisoprint/moonraker/error_info.xlsx\\nlanguage: en\\n'; "
                "with open(path, 'w') as f: f.write(content)"
            )
            # Escape for shell
            py_script = py_script.replace('"', '\\"')
            s, o = self.run_cmd(f"python3 -c \"{py_script}\"")
            if s:
                self.log("Config fixed. Restarting...")
                self.run_cmd(f"echo {self.password} | sudo -S systemctl restart moonraker")
                time.sleep(5)
                return True
            else:
                self.log(f"Failed to fix config: {o}")
                return False
        return True

    def load_test_cases_from_xlsx(self) -> List[str]:
        """
        从本地 error_info.xlsx 读取测试用例。
        使用第二列「报错信息」模板（msg_match）作为触发文本。
        """
        if not os.path.exists(self.xlsx_path):
            raise RuntimeError(f"xlsx not found: {self.xlsx_path}")
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise RuntimeError("Need openpyxl to read xlsx. pip install openpyxl")

        wb = load_workbook(self.xlsx_path, read_only=True, data_only=True)
        ws = wb.active
        if ws is None:
            wb.close()
            raise RuntimeError("xlsx has no active sheet")
        cases: List[str] = []
        header_skipped = False
        for row in ws.iter_rows(min_row=1, values_only=True):
            row = tuple(row) if row else ()
            if not row:
                continue
            first = str(row[0] or "").strip().lstrip("\ufeff")
            if not header_skipped:
                if "Error Code" in first or "Code" in first:
                    header_skipped = True
                    continue
                if first and not first.isdigit():
                    header_skipped = True
                    continue
            if len(row) < 2:
                continue
            msg = str(row[1] or "").strip()
            if msg:
                cases.append(msg)
        wb.close()
        seen = set()
        uniq: List[str] = []
        for c in cases:
            if c not in seen:
                seen.add(c)
                uniq.append(c)
        return uniq

    def trigger_errors_bulk(self, test_cases: List[str]) -> None:
        """
        一次性批量触发：把所有用例拼成多行脚本，只执行一次 printer.gcode.script。
        这比逐条 SSH+curl 快很多。
        """
        # 每条用例触发一行输出，保持和原逻辑一致
        script_lines = [f"M118 !! {case}" for case in test_cases]
        gcode_script = "\n".join(script_lines)
        payload = json.dumps({"script": gcode_script})
        cmd = (
            "curl -s -X POST "
            "-H 'Content-Type: application/json' "
            f"--data-raw {shlex.quote(payload)} "
            "http://localhost:7125/printer/gcode/script"
        )
        ok, out = self.run_cmd(cmd, check_exit=False, timeout=120)
        if not ok:
            self.log(f"Bulk trigger command failed: {out}")

    def check_logs_for_notification(self, error_msg, start_time):
        # We grep the log for "Error Indexer: Error detected" after our start time
        # Since we can't easily filter by time with grep, we just look at the last N lines
        # and see if we find our message.
        
        # But wait, we want to verify EACH error.
        # Let's count how many we found.
        pass

    def run_test_cases(self):
        test_cases = self.load_test_cases_from_xlsx()

        self.log(f"Starting test for {len(test_cases)} cases...")
        # Record current log size so we only parse new lines
        ok, out = self.run_cmd(f"stat -c %s {self.log_path}", check_exit=False)
        try:
            start_size = int(out.strip()) if ok and out.strip().isdigit() else 0
        except Exception:
            start_size = 0
        
        self.trigger_errors_bulk(test_cases)
        self.log("Triggered all cases in one request.")
        
        # Wait for Moonraker to process and log all notifications
        time.sleep(3)
        
        # 2. Check logs
        self.log("Verifying logs...")
        # Read only new log content since we started (avoid missing lines due to tail truncation)
        if start_size > 0:
            # tail -c +N is 1-based
            read_cmd = f"tail -c +$(( {start_size} + 1 )) {self.log_path}"
        else:
            # Fallback: read last chunk
            read_cmd = f"tail -n {len(test_cases) * 20} {self.log_path}"
        success, output = self.run_cmd(read_cmd, check_exit=False)
        
        if not success:
            self.log("Failed to read logs!")
            return

        detected_count = 0
        failed_cases = []
        
        # We look for "Error Indexer: Error detected: {code} - {msg}"
        # We need to map back to our test cases?
        # Let's just check if we have enough detections.
        
        # Count detections
        matches = re.findall(r"Error Indexer: Error detected: (\w+) -", output)
        detected_count = len(matches)

        # Identify missing cases by searching the error message part in new logs
        for case in test_cases:
            if case not in output:
                failed_cases.append(case)
        
        self.log(f"Triggered: {len(test_cases)}")
        self.log(f"Detected:  {detected_count}")
        
        if detected_count == len(test_cases):
            self.log("✅ All test cases passed!")
        else:
            self.log("❌ Some cases failed!")
            if failed_cases:
                self.log(f"Missing cases ({len(failed_cases)}):")
                for c in failed_cases:
                    self.log(f"  - {c}")
            unique_codes = set(matches)
            self.log(f"Unique codes found: {len(unique_codes)}")
            self.log(f"Codes found: {sorted(list(unique_codes))}")

if __name__ == "__main__":
    tester = MoonrakerTester("192.168.110.10", "anisoprint", "anaisuo")
    if tester.fix_config_if_needed():
        tester.run_test_cases()
