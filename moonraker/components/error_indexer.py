import logging
import re
import os
import time
from typing import Dict, Any, List, Optional, Tuple

from ..confighelper import ConfigHelper
from ..common import WebRequest
from ..utils import ServerError


DB_NAMESPACE = "moonraker"
WEBHOOK_ERROR_STATES = ("shutdown", "error")
PLACEHOLDER_RE = re.compile(r"%(\d+)")
DEFAULT_LANGUAGE = "en-US"
COMMON_SHEET_NAME = "common"


class ErrorIndexer:
    def __init__(self, config: ConfigHelper) -> None:
        self.server = config.get_server()
        self.name = config.get_name()
        # 获取 JSON-RPC 组件
        self.jsonrpc = self.server.lookup_component("jsonrpc")
        # 获取错误数据库文件路径
        self.file_path = self._get_file_path(config)
        # 默认读取 xlsx 中名为 "klipper" 的表格，可通过配置覆盖
        self.sheet_name = config.get("sheet_name", "klipper")
        # AI 错误单独使用 ai sheet，可通过配置覆盖
        self.ai_sheet_name = config.get("ai_sheet_name", "ai")
        # 语言设置默认为英文，将在 component_init 中从 fibretouch.language 读取
        self.language = DEFAULT_LANGUAGE
        # 错误数据库
        self.error_db: Dict[str, Dict[str, Any]] = {}
        # AI 错误数据库（来自单独的 ai sheet）
        self.ai_error_db: Dict[str, Dict[str, Any]] = {}
        # common 专有名词翻译表
        self.common_terms: Dict[str, Dict[str, Dict[str, str]]] = {}
        # 加载错误数据库
        self._load_error_db()
        self._load_ai_error_db()
        self._load_common_terms()
        logging.info(
            "Error Indexer: Component loaded (module=%s, file_path=%s, language=%s, db_size=%d)",
            __file__, self.file_path, self.language, len(self.error_db)
        )
        # 注册事件处理程序
        self._register_handlers()
        # 注册 WebSocket 通知
        self._register_notifications()
        # 挂钩 printer.gcode.script API
        self._hook_gcode_api()

    def _get_file_path(self, config: ConfigHelper) -> str:
        """xlsx path: prioritize xlsx_path, then file_path, csv_path, finally default error_info.xlsx."""
        default_path = os.path.abspath(
            os.path.join(os.path.dirname(__file__), "../../error_info.xlsx")
        )
        return str(config.get(
            "xlsx_path",
            config.get("file_path", config.get("csv_path", default_path)),
        ))

    # 注册事件处理程序
    def _register_handlers(self) -> None:
        # 1) 捕获 webhooks 的停机/错误信息
        self.server.register_event_handler(
            "server:status_update", self._handle_status_update
        )
        # 2) 捕获 Klipper 输出的 G-code 响应（含 "!!" 标准错误）
        self.server.register_event_handler(
            "server:gcode_response", self._handle_gcode_response
        )
        # 3) 捕获预检查失败（klippy_apis 在 _run_precheck 中发出）
        self.server.register_event_handler(
            "precheck:failed", self._handle_precheck_failed
        )
        # 4) 捕获 AI 检测结果事件（由 ai_detection 组件发送）
        self.server.register_event_handler(
            "ai_detection:detection_result", self._handle_ai_detection_result
        )


    # 注册 WebSocket 通知
    def _register_notifications(self) -> None:
        self.server.register_notification("error_indexer:error_detected")

    # 组件初始化
    async def component_init(self) -> None:
        """组件初始化：订阅 database:data_changed 事件并获取当前语言设置。"""
        # 注册 database:data_changed 事件处理，监听 fibretouch.language 变化
        self.server.register_event_handler(
            "database:data_changed",
            self._handle_database_data_changed
        )

        # 从数据库获取当前语言设置
        try:
            database = self.server.lookup_component("database", None)
            if database:
                language = await database.get_item("fibretouch", "language", None)
                if language:
                    normalized = self._normalize_locale(language)
                    if normalized:
                        self.language = normalized
                        logging.info("Error Indexer: Initial language setting from database: %s", self.language)
        except Exception as e:
            logging.warning(
                "Error Indexer: Failed to get initial language from database, using default %s: %s",
                self.language, e
            )

    def _handle_database_data_changed(self, event_data: Dict[str, Any]) -> None:
        """处理 database:data_changed 事件，更新语言设置。"""
        if not isinstance(event_data, dict):
            return

        # 检查是否是 fibretouch 命名空间的 language 字段变化
        namespace = event_data.get("namespace")
        records = event_data.get("records")

        if namespace == "fibretouch" and isinstance(records, dict) and "language" in records:
            new_language = records["language"]
            if new_language:
                normalized = self._normalize_locale(new_language)
                if normalized and normalized != self.language:
                    old_language = self.language
                    self.language = normalized
                    logging.info(
                        "Error Indexer: Language changed from %s to %s",
                        old_language, self.language
                    )

    def _hook_gcode_api(self) -> None:
        """Hook printer.gcode.script API to capture error responses."""
        method_info = self.jsonrpc.get_method("printer.gcode.script")
        if not method_info:
            return
        _, api_def = method_info
        orig_callback = api_def.callback
        async def wrapped_callback(web_request: WebRequest) -> Any:
            try:
                return await orig_callback(web_request)
            except ServerError as e:
                self._handle_gcode_response(f"!! {str(e)}")
                raise
        api_def.callback = wrapped_callback
        logging.info("Error Indexer: Hooked printer.gcode.script API")

    def _row_get(self, row: Tuple[Any, ...], idx: int) -> str:
        """Get the value from the row at index idx, convert to str and strip."""
        if idx < len(row):
            val = row[idx]
            return str(val).strip() if val is not None else ""
        return ""

    def _parse_placeholders(self, template: str) -> List[str]:
        """Parse placeholders from the template in the second column, e.g. ['%1','%2','%3']."""
        return list(dict.fromkeys(f"%{m.group(1)}" for m in PLACEHOLDER_RE.finditer(template)))

    def _build_msg_re(self, template: str, placeholders: List[str]) -> Optional[re.Pattern]:
        """Build regex for matching: %N appears as (.*?) first, then as a backreference. re.escape does not escape %, so escape %N segments separately."""
        if not template:
            return None
        if not placeholders:
            try:
                # 无占位符时也使用整行匹配，避免子串误命中
                pattern = "^" + re.escape(template) + "$"
                return re.compile(pattern, re.IGNORECASE)
            except re.error:
                return None
        parts = re.split(r"(%\d+)", template)
        group_for_ph: Dict[str, int] = {}
        next_group = 1
        result_parts: List[str] = []
        total_parts = len(parts)
        for idx, part in enumerate(parts):
            if PLACEHOLDER_RE.fullmatch(part):
                ph = part
                if ph in group_for_ph:
                    # 重复占位符使用反向引用，确保多次出现时值一致
                    result_parts.append("\\" + str(group_for_ph[ph]))
                else:
                    group_for_ph[ph] = next_group
                    # 针对「%N:%M」这类范围模式，限制前一个占位符不跨越冒号，
                    # 避免将整个区间都匹配到第一个占位符，导致第二个占位符为空
                    next_literal = parts[idx + 1] if idx + 1 < total_parts else ""
                    if next_literal and not PLACEHOLDER_RE.fullmatch(next_literal) \
                            and next_literal.startswith(":"):
                        # 形如 \"%3:%4\"，则 %3 仅匹配到冒号前的内容
                        result_parts.append("([^:]*)")
                    else:
                        result_parts.append("(.*?)")
                    next_group += 1
            else:
                result_parts.append(re.escape(part))
        try:
            # 使用全行匹配，避免只匹配到前半段导致尾部占位符为空
            pattern = "^" + "".join(result_parts) + "$"
            return re.compile(pattern, re.IGNORECASE)
        except re.error:
            logging.warning("Error Indexer: Invalid regex from template: %s", template[:80])
            return None

    def _load_error_db(self) -> None:
        """Load error table from xlsx specified sheet (skipping first row)."""
        if not os.path.exists(self.file_path):
            logging.warning("Error Indexer: xlsx file not found: %s", self.file_path)
            return
        try:
            from openpyxl import load_workbook
        except ImportError:
            logging.error("Error Indexer: Need openpyxl to read xlsx, please install openpyxl")
            return
        try:
            wb = load_workbook(self.file_path, read_only=True, data_only=True)
            if self.sheet_name not in wb.sheetnames:
                logging.warning(
                    "Error Indexer: xlsx sheet %r not found, available sheets: %s",
                    self.sheet_name, wb.sheetnames,
                )
                wb.close()
                return
            ws = wb[self.sheet_name]
            self.error_db = {}
            # 无论内容如何都先跳过第一行（表头）
            for row in ws.iter_rows(min_row=2, values_only=True):
                row = tuple(row) if row else ()
                if not row:
                    continue
                code = self._row_get(row, 0)
                if code.startswith("\ufeff"):
                    code = code[1:]
                if len(row) < 2:
                    continue
                if not code:
                    continue
                msg_match = self._row_get(row, 1)
                placeholders = self._parse_placeholders(msg_match)
                msg_re = self._build_msg_re(msg_match, placeholders)
                entry = {
                    "code": code,
                    "msg_match": msg_match,
                    "placeholders": placeholders,
                    "msg_re": msg_re,
                    "msg_multi": self._parse_multi_lang(self._row_get(row, 2)),
                    "solution_multi": self._parse_solution(self._row_get(row, 3)),
                }
                self.error_db[code] = entry
            wb.close()
            if self.error_db:
                logging.info(
                    "Error Indexer: Successfully loaded %d error records from xlsx sheet %s",
                    len(self.error_db),
                    self.sheet_name,
                )
        except Exception as e:
            logging.error("Error Indexer: Failed to load xlsx: %s", e)

    def _load_ai_error_db(self) -> None:
        """Load AI error table from dedicated ai sheet."""
        if not os.path.exists(self.file_path):
            return
        try:
            from openpyxl import load_workbook
        except ImportError:
            # 主库已经在 _load_error_db 中报过一次，这里安静失败即可
            return
        try:
            wb = load_workbook(self.file_path, read_only=True, data_only=True)
            if self.ai_sheet_name not in wb.sheetnames:
                wb.close()
                return
            ws = wb[self.ai_sheet_name]
            self.ai_error_db = {}
            for row in ws.iter_rows(min_row=2, values_only=True):
                row = tuple(row) if row else ()
                if not row:
                    continue
                code = self._row_get(row, 0)
                if code.startswith("\ufeff"):
                    code = code[1:]
                if len(row) < 2:
                    continue
                if not code:
                    continue
                msg_match = self._row_get(row, 1)
                placeholders = self._parse_placeholders(msg_match)
                msg_re = self._build_msg_re(msg_match, placeholders)
                entry = {
                    "code": code,
                    "msg_match": msg_match,
                    "placeholders": placeholders,
                    "msg_re": msg_re,
                    "msg_multi": self._parse_multi_lang(self._row_get(row, 2)),
                    "solution_multi": self._parse_solution(self._row_get(row, 3)),
                }
                self.ai_error_db[code] = entry
            wb.close()
            if self.ai_error_db:
                logging.info(
                    "Error Indexer: Successfully loaded %d AI error records from xlsx sheet %s",
                    len(self.ai_error_db),
                    self.ai_sheet_name,
                )
        except Exception as e:
            logging.error("Error Indexer: Failed to load AI xlsx sheet %s: %s", self.ai_sheet_name, e)

    def _load_common_terms(self) -> None:
        """从固定的 common sheet 加载占位符专有名词翻译。"""
        if not os.path.exists(self.file_path):
            return
        try:
            from openpyxl import load_workbook
        except ImportError:
            return
        try:
            wb = load_workbook(self.file_path, read_only=True, data_only=True)
            if COMMON_SHEET_NAME not in wb.sheetnames:
                logging.warning(
                    "Error Indexer: xlsx sheet %r not found for common terms",
                    COMMON_SHEET_NAME,
                )
                wb.close()
                return
            ws = wb[COMMON_SHEET_NAME]
            self.common_terms = {}
            for row in ws.iter_rows(min_row=2, values_only=True):
                row = tuple(row) if row else ()
                if not row:
                    continue
                key = self._row_get(row, 0)
                if key.startswith("\ufeff"):
                    key = key[1:]
                if not key:
                    continue
                self.common_terms[key] = {
                    "default": self._parse_multi_lang(self._row_get(row, 1)),
                    "10049": self._parse_multi_lang(self._row_get(row, 2)),
                }
            wb.close()
            if self.common_terms:
                logging.info(
                    "Error Indexer: Successfully loaded %d common terms from xlsx sheet %s",
                    len(self.common_terms),
                    COMMON_SHEET_NAME,
                )
        except Exception as e:
            logging.error(
                "Error Indexer: Failed to load common xlsx sheet %s: %s",
                COMMON_SHEET_NAME, e
            )

    def _parse_multi_lang(self, text: str) -> Dict[str, str]:
        """解析多语言字符串，支持 xx-XX: 或 xx-XX： 开头，每行一种语言。"""
        if not text:
            return {}
        text = text.replace("：", ":").replace("\r\n", "\n")
        # 按「语言行」切分，成对得到 (lang, content)
        parts = re.split(r"(?m)^([A-Za-z]{2}[-_][A-Za-z]{2}):", text)
        result = {}
        for i in range(1, len(parts), 2):
            if i + 1 < len(parts):
                lang = self._normalize_locale(parts[i].strip())
                if lang:
                    result[lang] = parts[i + 1].strip()
        return result

    def _normalize_locale(self, locale: str) -> str:
        """规范化为 xx-XX（小写语言+大写地区，如 zh-CN / en-US）。"""
        raw = (locale or "").strip()
        if not raw:
            return ""
        raw = raw.replace("_", "-")
        parts = [p for p in raw.split("-") if p]
        if len(parts) != 2:
            return ""
        language = parts[0].lower()
        region = parts[1].upper()
        if not re.fullmatch(r"[a-z]{2}", language):
            return ""
        if not re.fullmatch(r"[A-Z]{2}", region):
            return ""
        return f"{language}-{region}"

    def _locale_candidates(self, locale: str) -> List[str]:
        """返回候选 locale 列表（当前仅精确匹配，如 zh-CN）。"""
        norm = self._normalize_locale(locale)
        return [norm] if norm else []

    def _parse_solution(self, text: str) -> Dict[str, str]:
        """解析解决方法列：多语言（同 _parse_multi_lang）或单个 http(s) 链接。"""
        raw = (text or "").strip()
        if not raw:
            return {}
        if raw.startswith(("http://", "https://")):
            return {"en-US": raw}
        return self._parse_multi_lang(raw)

    def _match_error_code_in_db(
        self,
        message: str,
        db: Dict[str, Dict[str, Any]],
    ) -> Optional[Tuple[str, Tuple[str, ...]]]:
        """在指定错误库中匹配错误消息，返回 (code, 占位符捕获值) 或 None。"""
        for code, entry in db.items():
            msg_re = entry.get("msg_re")
            if msg_re is None:
                continue
            # 统一使用 fullmatch，避免正则 search 导致的子串误匹配
            m = msg_re.fullmatch(message)
            if not m:
                continue
            groups = m.groups() if m.lastindex else ()
            return (code, groups)
        return None

    def _match_error_code(self, message: str) -> Optional[Tuple[str, Tuple[str, ...]]]:
        """匹配主错误库（klipper sheet）中的错误消息。"""
        return self._match_error_code_in_db(message, self.error_db)

    def _resolve_placeholder_value(
        self, error_code: str, raw_value: str, lang: str
    ) -> str:
        """通过 common 词条解析单个占位符，未命中时保留原始值。"""
        entry = self.common_terms.get(raw_value)
        if not entry:
            return raw_value
        if error_code == "10049":
            special_multi = entry.get("10049", {}) or {}
            for cand in self._locale_candidates(lang):
                translated = special_multi.get(cand, "")
                if translated:
                    return translated
        default_multi = entry.get("default", {}) or {}
        for cand in self._locale_candidates(lang):
            translated = default_multi.get(cand, "")
            if translated:
                return translated
        return raw_value

    def _resolve_placeholder_values(
        self,
        error_code: str,
        placeholder_values: Tuple[str, ...],
        lang: str,
    ) -> Tuple[str, ...]:
        """在渲染本地化描述前，统一解析所有占位符。"""
        return tuple(
            self._resolve_placeholder_value(error_code, raw_value, lang)
            for raw_value in placeholder_values
        )

    def _handle_gcode_response(self, response: str) -> None:
        """Process G-code response, !!开头视为标准错误并匹配."""
        logging.debug(f"Error Indexer: Received response: {response}")

        clean_response = response.strip()
        clean_msg = clean_response[2:].strip() if clean_response.startswith("!!") else clean_response
        match_result = self._match_error_code(clean_msg)
        if match_result:
            code, groups = match_result
            self._process_error(clean_msg, response, code=code, placeholder_values=groups)
        elif clean_response.startswith("!!"):
            logging.info("Error Indexer: Standard error not found in database, processing as UNKNOWN: %s", clean_msg)
            self._process_error(clean_msg, response, None)

    async def _handle_status_update(self, status: Dict[str, Any]) -> None:
        """处理来自 notify_status_update 的状态：
        - webhooks.state 为 shutdown/error 时，使用 state_message 匹配并通知
        - print_stats.state 为 error 时，使用其中的 message 匹配并通知
        """
        # 1) 处理 webhooks 的关机/错误信息
        wh = status.get("webhooks")
        if isinstance(wh, dict):
            state = wh.get("state")
            if state in WEBHOOK_ERROR_STATES:
                state_message = wh.get("state_message", "")
                if state_message:
                    match_result = self._match_error_code(state_message)
                    original_response = (
                        f"webhooks.state={state}, state_message={state_message}"
                    )
                    if match_result:
                        code, groups = match_result
                        self._process_error(
                            state_message,
                            original_response,
                            code=code,
                            placeholder_values=groups,
                        )
                    else:
                        self._process_error(state_message, original_response, None)

        # 2) 处理 print_stats 的错误信息（例如预检查/预备阶段失败）
        ps = status.get("print_stats")
        if isinstance(ps, dict) and ps.get("state") == "error":
            message = ps.get("message", "") or ""
            if not message:
                return
            # 构造一个尽量接近文档示例的原始载荷，方便上位机调试
            origin_ps = {
                "filename": ps.get("filename"),
                "info": ps.get("info"),
                "message": message,
                "state": ps.get("state"),
                "total_duration": ps.get("total_duration"),
            }
            original_response = f"print_stats.error={origin_ps!r}"
            # 使用 message 走统一的错误码匹配和占位符替换流程
            match_result = self._match_error_code(message)
            if match_result:
                code, groups = match_result
                self._process_error(
                    message,
                    original_response,
                    code=code,
                    placeholder_values=groups,
                )
            else:
                self._process_error(message, original_response, None)

    def _handle_precheck_failed(self, payload: Dict[str, Any]) -> None:
        """处理预检查失败事件，由 klippy_apis._run_precheck 发出。"""
        message = (payload.get("message") or "").strip()
        if not message:
            return
        step = payload.get("step", "precheck")
        http_code = payload.get("http_code")
        data = payload.get("data")
        extra = {"step": step, "http_code": http_code}
        if data is not None:
            extra["data"] = data
        original_response = f"precheck:failed={extra!r}"
        match_result = self._match_error_code(message)
        if match_result:
            code, groups = match_result
            self._process_error(
                message,
                original_response,
                code=code,
                placeholder_values=groups,
            )
        else:
            self._process_error(message, original_response, None)

    def _handle_ai_detection_result(self, payload: Dict[str, Any]) -> None:
        """处理来自 ai_detection 组件的检测结果事件。

        该事件由 AIDetectionClient._send_detection_notification 发送，payload 示例:
            {
                "scenario": "printing_defect_periodic",
                "detection_type": "defect",
                "message": "Defect detected: spaghetti"
            }
        这里将其统一转换为 error_indexer:error_detected 通知，便于上位机展示。
        """
        if not isinstance(payload, dict):
            return

        message = (payload.get("message") or "").strip()
        if not message:
            return

        scenario = payload.get("scenario") or ""
        detection_type = payload.get("detection_type") or ""
        original_response = (
            f"ai_detection:detection_result "
            f"scenario={scenario}, detection_type={detection_type}, message={message}"
        )

        # 优先在 AI sheet 中进行错误码匹配
        match_result = self._match_error_code_in_db(message, self.ai_error_db) \
            if self.ai_error_db else None

        if match_result:
            code, groups = match_result
            self._process_error(
                message,
                original_response,
                code=code,
                placeholder_values=groups,
                error_db=self.ai_error_db,
            )
        else:
            # 未匹配到时仍然走统一流程，但使用 AI 错误库，
            # 最终会落到 UNKNOWN，description 为原始 message。
            self._process_error(
                message,
                original_response,
                error_db=self.ai_error_db,
            )

    def _process_error(
        self,
        error_msg: str,
        original_response: str,
        code: Optional[str] = None,
        placeholder_values: Optional[Tuple[str, ...]] = None,
        error_db: Optional[Dict[str, Dict[str, Any]]] = None,
    ) -> None:
        """Process detected error."""
        db = error_db or self.error_db
        if not code:
            # 按给定错误库匹配
            match_result = self._match_error_code_in_db(error_msg, db)
            if match_result:
                code, placeholder_values = match_result
        if not code:
            code = "UNKNOWN"
        if placeholder_values is None:
            placeholder_values = ()
        resolved_placeholder_values = self._resolve_placeholder_values(
            code, placeholder_values, self.language
        )

        entry = db.get(code, {})

        if code == "UNKNOWN":
            description = error_msg
            solution = ""
        else:
            description = self._get_description(
                entry, self.language, placeholder_values=resolved_placeholder_values
            )
            # 若当前语言没有对应文本，则按需求回退到原始错误信息
            if not description:
                description = error_msg
            solution = self._get_solution(entry, self.language)
        
        payload = {
            "error_code": code,
            "error_message": error_msg,
            "description": description,
            "timestamp": time.time(),
            "solution_link": solution,
            "original_response": original_response
        }

        self.server.send_event("error_indexer:error_detected", payload)
        logging.info(f"Error Indexer: Error detected and notification sent: {code} - {error_msg}")

    def _get_description(
        self,
        entry: Dict[str, Any],
        lang: str,
        placeholder_values: Optional[Tuple[str, ...]] = None,
    ) -> str:
        """获取错误描述。"""
        msg_multi = entry.get("msg_multi", {}) or {}
        for cand in self._locale_candidates(lang):
            template = msg_multi.get(cand, "")
            if not template:
                continue
            if placeholder_values and entry.get("placeholders"):
                placeholders = entry["placeholders"]
                filled = template
                for i, ph in enumerate(placeholders):
                    if i < len(placeholder_values) and ph in filled:
                        filled = filled.replace(ph, placeholder_values[i])
                return filled
            return template
        return ""

    def _get_solution(self, entry: Dict[str, Any], lang: str) -> str:
        """获取解决方法：当前语言 → en-US 兜底 → 任意非空。"""
        sol_multi = entry.get("solution_multi", {}) or {}
        for cand in self._locale_candidates(lang):
            if sol_multi.get(cand, ""):
                return sol_multi[cand]
        if sol_multi.get("en-US", ""):
            return sol_multi["en-US"]
        return next((v for v in sol_multi.values() if v), "")

def load_component(config: ConfigHelper) -> ErrorIndexer:
    return ErrorIndexer(config)
